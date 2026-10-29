"""PPTX -> DocIR 解析器。

每张幻灯片 = 一个 Page；文本框段落、表格单元格、备注页文本按形状顺序归一为 Block。

关键取舍：
1. 标题用占位符识别（placeholder idx == 0，即 PPT 内置的 Title / Center Title），
   命中即 BlockType.TITLE，其余文本框段落按缩进级别判为 LIST_ITEM / PARAGRAPH；
   不依赖形状名，因为本地化版本（如「标题 1」）名字会变。
2. 幻灯片尺寸由 presentation.slide_width/height（EMU）换算为 DocIR 的点（point）。
3. 形状坐标同样换算为点，写入 Block.bbox，便于后续版面相关处理；形状无坐标时置零。
4. 备注（notes_slide）单独收尾，标 meta["notes"]=True，避免与正文混淆。
5. row / col 统一 0-based。
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Emu

from ..core.config import Config
from ..core.docir import Block, BlockState, BlockType, DocIR, Origin, Page, Run, Style
from ..core.problems import DocumentParseError

EMU_PER_PT = 12700.0
_DEFAULT_BBOX = (0.0, 0.0, 0.0, 0.0)


class PptxParseError(DocumentParseError):
    """PPTX 解析失败。

    ADR-016：原 6 个 ``*ParseError`` 除名字外完全一致却各自演化，
    现统一继承 :class:`~linguaforge.core.problems.DocumentParseError`，
    仅保留 ``fmt="pptx"`` 作为格式标签，上层可一次捕获。
    """

    def __init__(self, message: str = "", **kw) -> None:
        super().__init__(message, fmt="pptx", **kw)


class PptxParser:
    """解析 .pptx。用法：

    with PptxParser(cfg) as p:
        p.open(path)
        for page in p.iter_pages():
            ...
    """

    extensions: set = {".pptx"}

    def __init__(self, cfg=None) -> None:
        self.cfg = cfg or Config()
        self.path: str = ""
        self._prs: Any | None = None
        self._page_count: int = 0
        self._title: str = ""

    # ---------------- 生命周期 ----------------
    def supports(self, path) -> bool:
        return Path(path).suffix.lower() in self.extensions

    def open(self, path) -> PptxParser:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(f"文件不存在: {p}")
        if not self.supports(p):
            raise PptxParseError(f"PptxParser 不支持扩展名 {p.suffix or '(无)'}，仅支持 .pptx")
        try:
            self._prs = Presentation(str(p))
        except Exception as e:
            raise PptxParseError(
                f"无法打开 PPTX：{p.name}（文件损坏或不是合法 OOXML 包）: {e}"
            ) from e
        self.path = str(p)
        self._page_count = len(self._prs.slides)
        self._title = self._detect_title()
        return self

    def close(self) -> None:
        self._prs = None
        self._page_count = 0

    def __enter__(self) -> PptxParser:
        return self

    def __exit__(self, *exc) -> None:
        self.close()

    # ---------------- 属性 ----------------
    @property
    def page_count(self) -> int:
        return self._page_count

    @property
    def title(self) -> str:
        return self._title or (Path(self.path).stem if self.path else "")

    # ---------------- 逐页解析 ----------------
    def _detect_title(self) -> str:
        """标题优先级：core properties -> 首页标题占位符 -> 文件名 stem。"""
        assert self._prs is not None
        try:
            t = (self._prs.core_properties.title or "").strip()
        except Exception:
            t = ""
        if t:
            return t
        try:
            slides = self._prs.slides
            if len(slides):
                for shape in slides[0].shapes:
                    if _is_title_shape(shape):
                        txt = _shape_text(shape).strip()
                        if txt:
                            return txt[:200]
        except Exception:
            pass
        return Path(self.path).stem

    def parse_page(self, index: int = 0) -> Page:
        assert self._prs is not None, "parser not opened"
        if not 0 <= index < self._page_count:
            raise IndexError(f"幻灯片序号越界: {index}（共 {self._page_count} 页）")
        slide = self._prs.slides[index]
        width = float(self._prs.slide_width or 0) / EMU_PER_PT or 720.0
        height = float(self._prs.slide_height or 0) / EMU_PER_PT or 540.0
        page = Page(index=index, width=width, height=height)
        page.meta["source_format"] = "pptx"

        order = 0
        table_seq = 0
        try:
            shapes = list(slide.shapes)
        except Exception:
            shapes = []

        for shape in shapes:
            bbox = _shape_bbox(shape)
            try:
                has_table = bool(shape.has_table)
            except Exception:
                has_table = False
            try:
                has_tf = bool(shape.has_text_frame)
            except Exception:
                has_tf = False

            if has_table:
                tid = f"p{index}_t{table_seq}"
                table_seq += 1
                for b in _table_blocks(shape.table, index, order, tid, bbox):
                    page.blocks.append(b)
                    order += 1
                continue
            if not has_tf:
                continue

            is_title = _is_title_shape(shape)
            for para in shape.text_frame.paragraphs:
                b = _para_block(para, index, order, bbox, is_title)
                if b is not None:
                    page.blocks.append(b)
                    order += 1

        # 备注：放在正文之后，单独标记
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text or ""
                notes = notes.strip()
                if notes:
                    b = Block(
                        type=BlockType.PARAGRAPH,
                        text=re.sub(r"[ \t]+", " ", notes),
                        style=Style(size=12.0),
                        page_index=index,
                        reading_order=order,
                        origin=Origin.TEXT_LAYER,
                        confidence=1.0,
                    )
                    b.meta["notes"] = True
                    b.state = BlockState.EXTRACTED if b.is_translatable else BlockState.SKIPPED
                    page.blocks.append(b)
                    order += 1
        except Exception:
            pass  # 备注缺失属正常情况，不影响正文解析

        page.meta["block_count"] = len(page.blocks)
        return page

    def iter_pages(self, start: int = 0, end: int | None = None) -> Iterator[Page]:
        stop = self._page_count if end is None else min(end, self._page_count)
        for i in range(max(0, start), stop):
            yield self.parse_page(i)

    def to_docir(self, max_pages: int | None = None) -> DocIR:
        doc = DocIR(source_path=self.path, source_format="pptx", title=self.title)
        for p in self.iter_pages(end=max_pages):
            doc.pages.append(p)
        return doc


# ---------------- 模块级工具 ----------------
def _is_title_shape(shape: Any) -> bool:
    """是否为标题占位符（idx==0）。非占位符形状直接返回 False。"""
    try:
        if not shape.is_placeholder:
            return False
        return shape.placeholder_format.idx == 0
    except Exception:
        return False


def _shape_text(shape: Any) -> str:
    try:
        return shape.text_frame.text or ""
    except Exception:
        return ""


def _shape_bbox(shape: Any) -> tuple:
    """形状坐标 EMU -> point。缺失坐标时返回零框。"""
    try:
        left = float(shape.left or 0) / EMU_PER_PT
        top = float(shape.top or 0) / EMU_PER_PT
        width = float(shape.width or 0) / EMU_PER_PT
        height = float(shape.height or 0) / EMU_PER_PT
        return (left, top, left + width, top + height)
    except Exception:
        return _DEFAULT_BBOX


def _para_block(
    para: Any, page_index: int, order: int, bbox: tuple, is_title: bool
) -> Block | None:
    """文本框段落 -> Block。空段落丢弃。"""
    runs = _runs_of(para)
    text = "".join(r.text for r in runs)
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return None

    level = 0
    try:
        level = int(para.level or 0)
    except (TypeError, ValueError):
        level = 0
    align = "left"
    try:
        al = para.alignment
        if al is not None:
            align = str(al).split(" ")[0].lower()
    except Exception:
        align = "left"

    style = Style(align=align)
    if runs:
        style.font = next((r.font for r in runs if r.font), "")
        style.size = max((r.size for r in runs if r.size), default=0.0)
        style.bold = any(r.bold for r in runs)
        style.italic = all(r.italic for r in runs)

    if is_title:
        btype, style.level = BlockType.TITLE, 1
    elif level > 0:
        btype = BlockType.LIST_ITEM
    else:
        btype = BlockType.PARAGRAPH

    b = Block(
        type=btype,
        bbox=bbox,
        text=text,
        runs=runs,
        style=style,
        page_index=page_index,
        reading_order=order,
        origin=Origin.TEXT_LAYER,
        confidence=1.0,
    )
    if btype is BlockType.LIST_ITEM:
        b.meta["list_level"] = level
    b.state = BlockState.EXTRACTED if b.is_translatable else BlockState.SKIPPED
    return b


def _table_blocks(
    table: Any, page_index: int, order: int, table_id: str, bbox: tuple
) -> list[Block]:
    """幻灯片表格 -> TABLE_CELL 块列表。"""
    out: list[Block] = []
    try:
        rows = list(table.rows)
    except Exception:
        return out
    for r, row in enumerate(rows):
        try:
            cells = list(row.cells)
        except Exception:
            continue
        for c, cell in enumerate(cells):
            text = re.sub(r"\s+", " ", (cell.text or "")).strip()
            if not text:
                continue
            b = Block(
                type=BlockType.TABLE_CELL,
                bbox=bbox,
                text=text,
                style=Style(size=12.0),
                page_index=page_index,
                reading_order=order + len(out),
                table_id=table_id,
                row=r,
                col=c,
                origin=Origin.TEXT_LAYER,
                confidence=1.0,
            )
            b.meta["header"] = r == 0
            b.state = BlockState.EXTRACTED if b.is_translatable else BlockState.SKIPPED
            out.append(b)
    return out


def _runs_of(para: Any) -> list[Run]:
    """段落 run -> DocIR Run（字号由 EMU Length 转 pt）。"""
    out: list[Run] = []
    try:
        pptx_runs = list(para.runs)
    except Exception:
        return out
    for r in pptx_runs:
        text = r.text or ""
        if not text:
            continue
        font = ""
        size = 0.0
        bold = False
        italic = False
        color = ""
        try:
            font = r.font.name or ""
        except Exception:
            font = ""
        try:
            if r.font.size is not None:
                size = float(Emu(r.font.size).pt)
        except Exception:
            size = 0.0
        try:
            bold = bool(r.font.bold)
            italic = bool(r.font.italic)
        except Exception:
            bold = italic = False
        try:
            rgb = r.font.color.rgb
            if rgb is not None:
                color = str(rgb)
        except (AttributeError, ValueError, TypeError):
            color = ""  # 主题色没有 rgb，属正常情况
        out.append(Run(text=text, font=font, size=size, bold=bold, italic=italic, color=color))
    return out

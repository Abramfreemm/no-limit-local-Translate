"""DocIR -> PPTX 输出。

每张幻灯片 = 一个 Page，页面上所有块按顺序放进同一个文本框（保持阅读顺序）。

接口：
    render_pages(state, pages, total, bilingual) -> str
        state 保存跨窗口的 Presentation（PPTX 需跨多次调用持续累积幻灯片，
        无法像文本那样拼字符串），返回本窗口的文本摘要。
    write_all(pages, out_path, *, bilingual, title) -> str
    finalize(state, out_path) -> str

关键取舍：
1. 统一用空白版式（slide_layouts[6]）+ 单个文本框，避免不同模板占位符数量不一致
   导致的 IndexError；标题靠字号/加粗体现而非占位符。
2. 表格无法用文本框表达，退化为「一行一个段落、单元格用 | 分隔」，保证信息不丢。
3. 无内容的页不生成幻灯片，避免导出一堆空白页。
4. 页眉 / 页脚 / 页码块不写入。
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from ..core.docir import Block, BlockType, Page

_SKIP = {BlockType.HEADER, BlockType.FOOTER, BlockType.PAGE_NUMBER}
_EMU_PER_PT = 12700
_BLANK_LAYOUT = 6

_SIZE_TITLE = Pt(28)
_SIZE_HEADING = Pt(20)
_SIZE_BODY = Pt(14)
_SIZE_NOTE = Pt(10)
_COLOR_GRAY = RGBColor(0x80, 0x80, 0x80)


def render_pages(
    state: dict,
    pages: Sequence[Page],
    total: int = 0,
    bilingual: bool = False,
) -> str:
    """把一个窗口的页面追加为幻灯片，返回本窗口的文本摘要。

    state 约定：
        state["prs"]   : Presentation（首次调用时创建）
        state["sized"] : 是否已按首页尺寸设置过画布
        state["title"] : 可选的文档标题（仅在首页无标题块时补一行）
    """
    prs = state.get("prs")
    if prs is None:
        prs = Presentation()
        state["prs"] = prs
    state["total"] = total

    digest: list[str] = []
    for page in pages:
        _add_slide(prs, page, bilingual, state, digest)
    return "\n".join(digest)


def finalize(state: dict, out_path) -> str:
    """保存 state 中的 Presentation，返回路径字符串。"""
    prs = state.get("prs")
    if prs is None:
        prs = Presentation()
        state["prs"] = prs
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return str(out)


def write_all(
    pages: Sequence[Page],
    out_path,
    *,
    bilingual: bool = False,
    title: str = "",
) -> str:
    """一次性写出完整 PPTX，返回输出路径字符串。"""
    state: dict[str, Any] = {}
    prs = Presentation()
    state["prs"] = prs
    if title:
        try:
            prs.core_properties.title = title
        except Exception:
            pass
        state["title"] = title
    render_pages(state, pages, len(pages), bilingual)
    return finalize(state, out_path)


# ---------------- 内部实现 ----------------
def _add_slide(
    prs: Any, page: Page, bilingual: bool, state: dict, digest: list[str]
) -> None:
    blocks = [b for b in page.blocks_sorted if b.type not in _SKIP]
    if not blocks:
        return

    if not state.get("sized"):
        _apply_size(prs, page)
        state["sized"] = True

    try:
        slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])
    except IndexError:
        slide = prs.slides.add_slide(prs.slide_layouts[0])

    margin = int(prs.slide_width * 0.05)
    box = slide.shapes.add_textbox(
        margin,
        margin,
        int(prs.slide_width - margin * 2),
        int(prs.slide_height - margin * 2),
    )
    tf = box.text_frame
    tf.word_wrap = True

    first = True
    # 首页无标题块时，用文档标题补一行
    pending_title = state.pop("title", "") if page.index == 0 else ""
    if pending_title and not any(b.type is BlockType.TITLE for b in blocks[:3]):
        _add_line(tf, pending_title, _SIZE_TITLE, True, first)
        digest.append(pending_title)
        first = False

    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b.type is BlockType.TABLE_CELL:
            tid = b.table_id or ""
            group: list[Block] = []
            while (
                i < len(blocks)
                and blocks[i].type is BlockType.TABLE_CELL
                and (blocks[i].table_id or "") == tid
            ):
                group.append(blocks[i])
                i += 1
            for line in _digest_table(group):
                _add_line(tf, line, _SIZE_BODY, False, first)
                digest.append(line)
                first = False
            continue

        text = (b.out_text or "").strip()
        if not text:
            i += 1
            continue
        size, bold, level = _SIZE_BODY, False, 0
        if b.type is BlockType.TITLE:
            size, bold = _SIZE_TITLE, True
        elif b.type is BlockType.HEADING:
            size, bold = _SIZE_HEADING, True
        elif b.type is BlockType.LIST_ITEM:
            level = 1
        _add_line(tf, text, size, bold, first, level)
        digest.append(text)
        first = False

        if bilingual and b.translated.strip() and b.text.strip():
            _add_line(tf, b.text.strip(), _SIZE_NOTE, False, False, level, italic=True)
        i += 1


def _apply_size(prs: Any, page: Page) -> None:
    """按首页尺寸设置画布（DocIR 用点，PPTX 用 EMU）。"""
    try:
        if page.width and page.width > 10:
            prs.slide_width = Emu(int(round(page.width * _EMU_PER_PT)))
        if page.height and page.height > 10:
            prs.slide_height = Emu(int(round(page.height * _EMU_PER_PT)))
    except Exception:
        pass


def _add_line(
    tf: Any,
    text: str,
    size: Any,
    bold: bool,
    first: bool,
    level: int = 0,
    italic: bool = False,
) -> None:
    """往文本框追加一个段落。first=True 时复用自带空段落，避免留下空行。"""
    para = tf.paragraphs[0] if first else tf.add_paragraph()
    para.level = level
    run = para.add_run()
    run.text = text
    try:
        run.font.size = size
        run.font.bold = bold
        run.font.italic = italic
        if not bold and not italic and size == _SIZE_NOTE:
            run.font.color.rgb = _COLOR_GRAY
    except Exception:
        pass


def _digest_table(cells: Sequence[Block]) -> list[str]:
    """表格 -> 每行一行文本（单元格用 | 分隔）。"""
    rows: dict[int, dict[int, Block]] = {}
    for c in cells:
        rows.setdefault(c.row, {})[c.col] = c
    out: list[str] = []
    for ri in sorted(rows):
        row = rows[ri]
        line = " | ".join(row[c].out_text.strip() for c in sorted(row))
        if line.strip(" |"):
            out.append(line)
    return out

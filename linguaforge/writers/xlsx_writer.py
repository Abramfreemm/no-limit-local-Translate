"""DocIR -> XLSX 输出。

按 block.row / block.col 把单元格还原到工作表；一个 Page = 一个工作表。

接口：
    render_pages(state, pages, total, bilingual) -> str
        state 保存跨窗口的 Workbook，返回本窗口的文本摘要。
    write_all(pages, out_path, *, bilingual, title) -> str
    finalize(state, out_path) -> str

关键取舍：
1. DocIR 的 row / col 是 0-based，openpyxl 是 1-based，写回时统一 +1。
2. 非 TABLE_CELL 块（来自 PDF / DOCX 正文）没有坐标，退化为 A 列顺次写入，
   并接在表格下方（表格最大行 + 2），避免覆盖已有单元格。
3. 双语模式不改列结构（会破坏坐标），改为在同页后追加一个「<表名>_原文」工作表，
   按相同坐标放原文，保证无损且可读。
4. 全空 Page 仍生成空表，保持 Page 与工作表一一对应，便于核对。
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..core.docir import Block, BlockType, Page

_SKIP = {BlockType.HEADER, BlockType.FOOTER, BlockType.PAGE_NUMBER}
_MAX_SHEET_NAME = 31
_MAX_COL_WIDTH = 60.0
_INVALID_CHARS = '[]:*?/\\'


def render_pages(
    state: dict,
    pages: Sequence[Page],
    total: int = 0,
    bilingual: bool = False,
) -> str:
    """把一个窗口的页面追加为工作表，返回本窗口的文本摘要。

    state 约定：
        state["wb"]        : Workbook（首次调用时创建）
        state["names"]     : 已用工作表名（保证唯一）
        state["sheet_seq"] : 自增序号，用于生成默认表名
        state["prefer"]    : 单页时优先使用的表名（通常来自 title）
    """
    wb = state.get("wb")
    if wb is None:
        wb = Workbook()
        wb.remove(wb.active)
        state["wb"] = wb
        state["names"] = set()
        state["sheet_seq"] = 0
    state["total"] = total

    digest: list[str] = []
    for page in pages:
        _add_sheet(wb, page, bilingual, state, digest)
    return "\n".join(digest)


def finalize(state: dict, out_path) -> str:
    """保存 state 中的 Workbook，返回路径字符串。"""
    wb = state.get("wb")
    if wb is None:
        wb = Workbook()
        state["wb"] = wb
    if not wb.sheetnames:
        wb.create_sheet(title="Sheet1")  # 空工作簿无法保存
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out)


def write_all(
    pages: Sequence[Page],
    out_path,
    *,
    bilingual: bool = False,
    title: str = "",
) -> str:
    """一次性写出完整 XLSX，返回输出路径字符串。"""
    state: dict[str, Any] = {}
    wb = Workbook()
    wb.remove(wb.active)
    state["wb"] = wb
    if title:
        try:
            wb.properties.title = title
        except Exception:
            pass
        if len(pages) == 1:
            state["prefer"] = title
    render_pages(state, pages, len(pages), bilingual)
    return finalize(state, out_path)


# ---------------- 内部实现 ----------------
def _add_sheet(wb: Any, page: Page, bilingual: bool, state: dict, digest: list[str]) -> None:
    name = _next_sheet_name(state, page)
    ws = wb.create_sheet(title=name)
    ws.sheet_view.showGridLines = True

    blocks = [b for b in page.blocks_sorted if b.type not in _SKIP]
    cells = [b for b in blocks if b.type is BlockType.TABLE_CELL]
    others = [b for b in blocks if b.type is not BlockType.TABLE_CELL]

    src_ws = None
    if bilingual:
        src_ws = wb.create_sheet(title=_next_name(state, f"{name}_原文"))

    bold_font = Font(bold=True)
    for b in cells:
        text = b.out_text.strip()
        if not text:
            continue
        row = max(0, int(b.row)) + 1
        col = max(0, int(b.col)) + 1
        c = ws.cell(row=row, column=col, value=text)
        if b.meta.get("header") or b.style.bold:
            c.font = bold_font
        if src_ws is not None and b.text.strip():
            src_ws.cell(row=row, column=col, value=b.text.strip())

    # 无坐标的正文块：A 列顺次写，接在表格下方
    if others:
        start_row = (max(int(b.row) for b in cells) + 3) if cells else 1
        for i, b in enumerate(others):
            text = b.out_text.strip()
            if not text:
                continue
            row = start_row + i
            c = ws.cell(row=row, column=1, value=text)
            if b.type in (BlockType.TITLE, BlockType.HEADING) or b.style.bold:
                c.font = bold_font
            if src_ws is not None:
                src_ws.cell(row=row, column=1, value=b.text.strip())

    _autosize(ws)
    if src_ws is not None:
        _autosize(src_ws)

    digest.extend(_digest(cells, others))


def _digest(cells: Sequence[Block], others: Sequence[Block]) -> list[str]:
    rows: dict[int, dict[int, Block]] = {}
    for c in cells:
        rows.setdefault(c.row, {})[c.col] = c
    out: list[str] = []
    for ri in sorted(rows):
        row = rows[ri]
        out.append(" | ".join(row[c].out_text.strip() for c in sorted(row)))
    out.extend(b.out_text.strip() for b in others if b.out_text.strip())
    return out


def _next_sheet_name(state: dict, page: Page) -> str:
    """工作表名优先级：解析时记录的原始表名 > title > Sheet{N}。"""
    raw = str(page.meta.get("sheet_name") or "").strip()
    if not raw:
        raw = str(state.pop("prefer", "") or "").strip()
    if not raw:
        state["sheet_seq"] = int(state.get("sheet_seq", 0)) + 1
        raw = f"Sheet{state['sheet_seq']}"
    return _next_name(state, raw)


def _next_name(state: dict, raw: str) -> str:
    """清洗并去重工作表名（<=31 字符、不含 []:*?/\\）。"""
    used = state.setdefault("names", set())
    base = "".join(ch for ch in raw if ch not in _INVALID_CHARS).strip() or "Sheet"
    base = base[:_MAX_SHEET_NAME]
    name = base
    i = 2
    while name in used:
        suffix = f"_{i}"
        name = f"{base[: _MAX_SHEET_NAME - len(suffix)]}{suffix}"
        i += 1
    used.add(name)
    return name


def _display_width(s: str) -> float:
    """粗略显示宽度：中日韩字符按 2 个半角计。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in s)


def _autosize(ws: Any) -> None:
    """按内容设置列宽（上限 60），避免导出后所有列挤在一起。"""
    widths: dict[int, float] = {}
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            w = _display_width(str(value))
            if w > widths.get(cell.column, 0.0):
                widths[cell.column] = w
    for col, w in widths.items():
        try:
            ws.column_dimensions[get_column_letter(col)].width = min(
                _MAX_COL_WIDTH, max(8.0, w + 2.0)
            )
        except Exception:
            continue

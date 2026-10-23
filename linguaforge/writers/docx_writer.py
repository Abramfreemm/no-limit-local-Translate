"""DocIR -> DOCX 输出。

保留标题层级（Title / Heading 1-9）、段落、列表、表格；run 级加粗 / 斜体 / 字体尽量还原。

接口（与流水线窗口式调用对齐）：
    render_pages(state, pages, total, bilingual) -> str
        state 保存跨窗口的 python-docx Document（二进制格式无法返回文件内容，
        这里返回本窗口的文本摘要，供日志与预览使用）。
    write_all(pages, out_path, *, bilingual, title) -> str
        一次性写完整文件，返回输出路径。
    finalize(state, out_path) -> str
        流式场景收尾保存，返回输出路径。

关键取舍：
1. 块级样式优先用 Word 内置样式名（Title / Heading N / List Bullet），
   这样生成的文件在 Word / WPS 里大纲级别正确；样式缺失时回退为手工字号+粗体。
2. 只在 run 文本与块文本「去空白后一致」时逐 run 还原格式，否则整块用块级样式——
   翻译后字数变化是正常的，强行对齐 run 会错位。
3. 双语模式：译文为主段，原文紧随其后为小号斜体灰字段落（与 Markdown 的 > 引用对齐）。
4. 页眉 / 页脚 / 页码块不写入（这些是 PDF 版面产物，写进 DOCX 会变正文噪音）。
"""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..core.docir import Block, BlockType, Page

_SKIP = {BlockType.HEADER, BlockType.FOOTER, BlockType.PAGE_NUMBER}

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_HEADING_SIZES = {1: 20.0, 2: 16.0, 3: 14.0, 4: 13.0, 5: 12.0, 6: 11.5}


def render_pages(
    state: dict,
    pages: Sequence[Page],
    total: int = 0,
    bilingual: bool = False,
) -> str:
    """把一个窗口的页面追加进 state 里的 Document，返回本窗口的文本摘要。

    state 约定：
        state["doc"]  : python-docx Document（首次调用时创建）
        state["total"]: 总页数（仅记录）
        state["pages_done"]: 已写页数
    """
    doc = state.get("doc")
    if doc is None:
        doc = Document()
        state["doc"] = doc
    state["total"] = total
    state["pages_done"] = int(state.get("pages_done", 0))

    digest: list[str] = []
    for page in pages:
        _render_page(doc, page, bilingual, digest)
        state["pages_done"] += 1
    return "\n".join(digest)


def finalize(state: dict, out_path) -> str:
    """保存 state 中的 Document 到 out_path，返回路径字符串。"""
    doc = state.get("doc")
    if doc is None:
        doc = Document()
        state["doc"] = doc
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return str(out)


def write_all(
    pages: Sequence[Page],
    out_path,
    *,
    bilingual: bool = False,
    title: str = "",
) -> str:
    """一次性写出完整 DOCX，返回输出路径字符串。"""
    state: dict[str, Any] = {}
    doc = Document()
    state["doc"] = doc
    if title:
        try:
            doc.core_properties.title = title
        except Exception:
            pass
        # 正文已有标题块（或首个一级标题就等同于该标题）时不重复写，避免两份标题
        has_title = any(
            b.type is BlockType.TITLE
            or (
                b.type is BlockType.HEADING
                and (b.style.level or 2) <= 1
                and b.text.strip() == title.strip()
            )
            for p in pages
            for b in p.blocks_sorted[:5]
        )
        if not has_title:
            p = doc.add_paragraph(style="Title")
            p.add_run(title)
    render_pages(state, pages, len(pages), bilingual)
    return finalize(state, out_path)


# ---------------- 内部实现 ----------------
def _render_page(doc: Any, page: Page, bilingual: bool, digest: list[str]) -> None:
    blocks = [b for b in page.blocks_sorted if b.type not in _SKIP]
    i = 0
    while i < len(blocks):
        b = blocks[i]
        if b.type is BlockType.TABLE_CELL:
            tid = b.table_id or ""
            group: list[Block] = []
            while (
                i < len(blocks)
                and blocks[i].type is BlockType.TABLE_CELL
                and (blocks[i].table_id or "") == tid
            ):
                group.append(blocks[i])
                i += 1
            _add_table(doc, group)
            digest.extend(_digest_cells(group))
            continue
        _add_paragraph(doc, b, bilingual, digest)
        i += 1


def _digest_cells(cells: Sequence[Block]) -> list[str]:
    rows: dict[int, dict[int, Block]] = {}
    for c in cells:
        rows.setdefault(c.row, {})[c.col] = c
    out: list[str] = []
    for ri in sorted(rows):
        row = rows[ri]
        out.append(" | ".join(row[c].out_text.strip() for c in sorted(row)))
    return out


def _add_paragraph(doc: Any, b: Block, bilingual: bool, digest: list[str]) -> None:
    text = (b.out_text or "").strip()
    if not text:
        return

    style_name: str | None = None
    fallback_size = 11.0
    fallback_bold = b.style.bold
    if b.type is BlockType.TITLE:
        style_name, fallback_size, fallback_bold = "Title", 20.0, True
    elif b.type is BlockType.HEADING:
        level = max(1, min(9, b.style.level or 2))
        style_name = f"Heading {level}"
        fallback_size = _HEADING_SIZES.get(level, 11.5)
        fallback_bold = True
    elif b.type is BlockType.LIST_ITEM:
        style_name = "List Bullet"
    elif b.type is BlockType.CODE:
        style_name = None
        fallback_size = 10.5

    p = None
    if style_name:
        try:
            p = doc.add_paragraph(style=style_name)
        except KeyError:
            p = None  # 模板缺少该内置样式 -> 回退手工格式
    if p is None:
        p = doc.add_paragraph()
        if b.type is BlockType.LIST_ITEM:
            text = f"• {text}"

    try:
        if b.style.align in _ALIGN_MAP:
            p.alignment = _ALIGN_MAP[b.style.align]
    except Exception:
        pass

    _fill_runs(p, b, text, fallback_size, fallback_bold, b.style.italic)
    digest.append(text)

    # 双语：原文另起一段，小号斜体灰字
    if bilingual and b.translated.strip() and b.text.strip():
        orig = b.text.strip()
        p2 = doc.add_paragraph()
        r = p2.add_run(orig)
        r.italic = True
        r.font.size = Pt(max(8.0, fallback_size - 2.0))
        try:
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        except Exception:
            pass


def _fill_runs(
    p: Any,
    b: Block,
    text: str,
    size: float,
    bold: bool,
    italic: bool,
) -> None:
    """写入 run：能对齐就逐 run 还原格式，否则整块套用块级样式。"""
    src = b.runs or []
    joined = "".join(r.text for r in src)
    if src and _skeleton(joined) == _skeleton(text):
        base_size = size if size else 11.0
        for r in src:
            if not r.text:
                continue
            run = p.add_run(r.text)
            run.bold = bool(r.bold)
            run.italic = bool(r.italic)
            if r.font:
                _set_font(run, r.font)
            try:
                run.font.size = Pt(r.size if r.size else base_size)
            except Exception:
                pass
            if r.color:
                try:
                    run.font.color.rgb = RGBColor.from_string(r.color)
                except (ValueError, AttributeError):
                    pass
        return

    run = p.add_run(text)
    run.bold = bool(bold)
    run.italic = bool(italic)
    if b.style.font:
        _set_font(run, b.style.font)
    try:
        run.font.size = Pt(size if size else 11.0)
    except Exception:
        pass


def _skeleton(s: str) -> str:
    """去全部空白后比较，用于判断 run 序列是否仍与目标文本同构。"""
    return re.sub(r"\s+", "", s or "")


def _set_font(run: Any, name: str) -> None:
    """设置字体名，同时写 eastAsia，否则 Word 打开中文会回退到主题字体。"""
    try:
        run.font.name = name
        rPr = run._element.get_or_add_rPr()  # noqa: SLF001
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rFonts.set(qn(attr), name)
    except Exception:
        pass


def _add_table(doc: Any, cells: Sequence[Block]) -> None:
    """按 row / col 还原为真正的 DOCX 表格。"""
    if not cells:
        return
    rows: dict[int, dict[int, Block]] = {}
    for c in cells:
        rows.setdefault(c.row, {})[c.col] = c
    nrows = max(rows) + 1
    ncols = max((max(r) for r in rows.values() if r), default=0) + 1
    try:
        table = doc.add_table(rows=nrows, cols=ncols)
    except Exception:
        return
    try:
        table.style = "Table Grid"
    except KeyError:
        pass  # 模板无此样式也能正常打开，只是没有边框

    for ri in sorted(rows):
        row_cells = rows[ri]
        for ci in sorted(row_cells):
            b = row_cells[ci]
            try:
                cell = table.cell(ri, ci)
            except (IndexError, ValueError):
                continue
            cell.text = ""
            para = cell.paragraphs[0]
            text = b.out_text.strip()
            run = para.add_run(text)
            run.bold = bool(b.style.bold or b.meta.get("header"))
            run.italic = bool(b.style.italic)
            if b.style.font:
                _set_font(run, b.style.font)
